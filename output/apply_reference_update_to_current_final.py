from datetime import datetime
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"E:\Project_Metroidvania\Metroidvania")
SRC = ROOT / "output" / "current_final_reference_work.docx"
OUT = ROOT / "output" / "current_final_reference_updated.docx"
FINAL = ROOT / "葛广一" / "毕业设计论文最终版.docx"
SAVE_DIR = ROOT / "葛广一" / "save"
SAVE_DIR.mkdir(parents=True, exist_ok=True)
BACKUP = SAVE_DIR / f"毕业设计论文最终版_before_reference_update_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=12, bold=False, superscript=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.superscript = superscript
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.underline = False
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = rpr._new_rFonts()
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def clear_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = ""


def set_body_paragraph_with_citation(paragraph, parts):
    clear_paragraph(paragraph)
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Pt(24)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.25
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    for text, is_citation in parts:
        run = paragraph.add_run(text)
        if is_citation:
            set_run_font(run, "宋体", "Times New Roman", 10.5, False, True)
        else:
            set_run_font(run, "宋体", "Times New Roman", 12, False, False)


def set_reference_text(paragraph, text):
    clear_paragraph(paragraph)
    paragraph.style = next(s for s in paragraph.part.document.styles if s.style_id == "afa")
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.left_indent = Pt(21)
    fmt.first_line_indent = Pt(-21)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(17)
    fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", 10.5, False, False)


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style is not None:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def find_para(doc, predicate, label):
    for idx, paragraph in enumerate(doc.paragraphs):
        if predicate(paragraph.text.strip()):
            return idx, paragraph
    raise RuntimeError(f"Cannot find paragraph: {label}")


copy2(FINAL, BACKUP)
doc = Document(SRC)

_, p51 = find_para(doc, lambda t: t.startswith("Unity是一款2D和3D游戏开发时常用的工具"), "Unity overview")
set_body_paragraph_with_citation(
    p51,
    [
        (
            "Unity是一款2D和3D游戏开发时常用的工具，提供了Tilemap、Animator、Rigidbody2D、Collider2D、AudioSource、Canvas等组件和功能。相关资料中既涉及2D平台游戏实现、跳跃手感、游戏AI和游戏引擎架构，也包括编程模式、游戏体验设计以及Unity官方手册和脚本API等内容",
            False,
        ),
        ("[1-9]", True),
        (
            "，这些内容可以为本项目的玩家控制、场景搭建、动画切换和敌人行为设计提供参考。",
            False,
        ),
    ],
)

_, p52 = find_para(doc, lambda t: t.startswith("在游戏AI的设计中"), "AI paragraph")
set_body_paragraph_with_citation(
    p52,
    [
        (
            "在游戏AI的设计中，有限状态机、行为树、分层状态机等都是比较常见的组织方式，拿毕业设计规模的2D动作游戏Demo来说，有限状态机的实现成本低、结构直观、调试时容易看出对象当前处在哪个阶段；普通敌人用待机、巡逻、追踪、攻击、受击、死亡等状态来描述行为，Boss可以在普通敌人基础上加上传送、施法、二阶段转场等状态来表现更复杂的战斗节奏；与把全部逻辑堆在一个Update函数中的方式相比，状态机可以让每一个状态只负责本阶段要做的事情，后续维护会轻一些",
            False,
        ),
        ("[10-11]", True),
        ("。", False),
    ],
)

_, p53 = find_para(doc, lambda t: t.startswith("手感一般同输入响应速度"), "game feel paragraph")
set_body_paragraph_with_citation(
    p53,
    [
        (
            "手感一般同输入响应速度、移动加速度、跳跃高度、碰撞检测稳定度、攻击前后摇、受击反馈等有着联系，本项目并没有试图构建起一个商业作品复杂的程度，而把重心放在了毕业Demo可以展现的核心体验部分上，即玩家能够较为流畅地移动、跳跃、冲刺、攻击，敌人可以在一定的范围内找到玩家并做出反应，Boss会给出阶段变化以及明显的攻击提示；这样的取舍有利于在较短的开发周期内得到稳定的、可以答辩的系统成果",
            False,
        ),
        ("[12-14]", True),
        ("。", False),
    ],
)

_, p98 = find_para(doc, lambda t: t.startswith("本设计使用Unity引擎和C#语言进行开发"), "dev env paragraph")
set_body_paragraph_with_citation(
    p98,
    [
        ("本设计使用Unity引擎和C#语言进行开发，其主要的开发环境和工具见表3.1", False),
        ("[8-9,15]", True),
        ("。", False),
    ],
)

_, p101 = find_para(doc, lambda t: t.startswith("在技术选择上，Unity可以采用组件化"), "tech choice paragraph")
set_body_paragraph_with_citation(
    p101,
    [
        (
            "在技术选择上，Unity可以采用组件化的组织方式来管理GameObject和脚本，非常适合毕业设计中快速搭建可运行的作品，C#脚本用来描述各个模块的行为，对玩家状态切换、敌人检测、Boss攻击、技能生成、血量更新、菜单响应等都可以通过脚本和Unity组件一起实现；与从底层图形库开始实现相比，使用Unity可以让开发的重点更多地放在玩法逻辑和软件结构上，更符合本课题的目标",
            False,
        ),
        ("[8-9,16-17]", True),
        ("。", False),
    ],
)

_, p217 = find_para(doc, lambda t: t.startswith("Boss模块以Enemy_DeathBringer为核心"), "boss module paragraph")
base_217 = (
    "Boss模块以Enemy_DeathBringer为核心，该类继承Enemy，并在Awake中创建idleState、battleState、attackState、spellCastState、phaseTransitionState、teleportState和deadState。Boss默认处于idleState，当玩家进入bossFightDistance范围后，ShouldStartBossFight返回true，Boss才会开始进入战斗流程。和普通敌人相比，Boss拥有更多状态，切换条件也更复杂，有限状态机在这里能把较复杂的行为拆成可说明、可维护的几个阶段"
)
set_body_paragraph_with_citation(p217, [(base_217, False), ("[3,10-11]", True), ("。", False)])

# 参考文献标题格式保险处理。
_, ref_title = find_para(doc, lambda t: t == "参 考 文 献", "reference title")
ref_title.paragraph_format.page_break_before = True
ref_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
ref_title.paragraph_format.line_spacing = 1.5
ref_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
for run in ref_title.runs:
    set_run_font(run, "黑体", "Times New Roman", 15, True, False)

# 修正第 2 条会议论文，并新增两条 Unity/C# 开发相关文献。
refs = [(idx, p) for idx, p in enumerate(doc.paragraphs) if p.text.strip().startswith("[")]
ref_by_number = {}
for _, paragraph in refs:
    text = paragraph.text.strip()
    if text.startswith("["):
        num = int(text.split("]", 1)[0][1:])
        ref_by_number[num] = paragraph

set_reference_text(
    ref_by_number[2],
    "[2] Fasterholdt M, Pichlmair M, Holmgård C. You Say Jump, I Say How High? Operationalising the Game Feel of Jumping[C]//Proceedings of DiGRA/FDG 2016 Conference. DiGRA, 2016.",
)

last_ref = ref_by_number[max(ref_by_number)]
new16 = insert_paragraph_after(last_ref, style=last_ref.style)
set_reference_text(
    new16,
    "[16] Hocking J. Unity in Action: Multiplatform Game Development in C# with Unity 2022[M]. 3rd ed. Shelter Island: Manning Publications, 2022.",
)
new17 = insert_paragraph_after(new16, style=last_ref.style)
set_reference_text(
    new17,
    "[17] Bond J G. Introduction to Game Design, Prototyping, and Development: From Concept to Playable Game with Unity and C#[M]. 2nd ed. Boston: Addison-Wesley Professional, 2018.",
)

doc.save(OUT)
copy2(OUT, FINAL)
print(f"backup={BACKUP}")
print(f"out={OUT}")
print(f"final={FINAL}")
