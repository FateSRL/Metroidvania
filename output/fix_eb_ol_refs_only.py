from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SRC = "output/verify_refs_only_current.docx"
OUT = "output/verify_refs_only_current_fixed.docx"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.superscript = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = rpr._new_rFonts()
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def set_reference_text(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.left_indent = Pt(21)
    fmt.first_line_indent = Pt(-21)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(17)
    fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    run = paragraph.add_run(text)
    set_run_font(run)


doc = Document(SRC)
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text.startswith("[8] "):
        set_reference_text(paragraph, "[8] Unity Technologies. Unity User Manual 2022.3 LTS[EB/OL]. [2026-05-13]. https://docs.unity3d.com/2022.3/Documentation/Manual/index.html.")
    elif text.startswith("[9] "):
        set_reference_text(paragraph, "[9] Unity Technologies. Unity Scripting API 2022.3 LTS[EB/OL]. [2026-05-13]. https://docs.unity3d.com/2022.3/Documentation/ScriptReference/.")
    elif text.startswith("[15] "):
        set_reference_text(paragraph, "[15] Microsoft. C# documentation[EB/OL]. [2026-05-13]. https://learn.microsoft.com/dotnet/csharp/.")
doc.save(OUT)
print(OUT)
