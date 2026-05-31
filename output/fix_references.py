from pathlib import Path
from shutil import copy2
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"E:\Project_Metroidvania\Metroidvania")
SRC = ROOT / "葛广一" / "毕业设计论文最终版_格式调整版.docx"
OUT = ROOT / "葛广一" / "毕业设计论文最终版_参考文献修正版.docx"
SAVE_DIR = ROOT / "葛广一" / "save"
SAVE_DIR.mkdir(parents=True, exist_ok=True)
BACKUP = SAVE_DIR / f"毕业设计论文最终版_格式调整版_before_reference_fix_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=12, bold=False, superscript=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.superscript = superscript
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = rpr._new_rFonts()
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def clear_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = ""


def set_body_paragraph_with_citation(paragraph, parts):
    """parts is list of (text, is_citation)."""
    clear_paragraph(paragraph)
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Pt(24)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.25
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    for text, is_citation in parts:
        run = paragraph.add_run(text)
        if is_citation:
            set_run_font(run, "宋体", "Times New Roman", 10.5, False, True)
        else:
            set_run_font(run, "宋体", "Times New Roman", 12, False, False)


def set_reference_text(paragraph, text):
    clear_paragraph(paragraph)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Pt(21)
    paragraph.paragraph_format.first_line_indent = Pt(-21)
    paragraph.paragraph_format.line_spacing = Pt(17)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", 10.5, False, False)


copy2(SRC, BACKUP)
doc = Document(SRC)

# 1.2 国内外研究现状：用顺序编码第一次覆盖 1-9。
p = doc.paragraphs[51]
set_body_paragraph_with_citation(
    p,
    [
        (
            "Unity是一款2D和3D游戏开发时常用的工具，提供了Tilemap、Animator、Rigidbody2D、Collider2D、AudioSource、Canvas等组件和功能。相关资料中既涉及2D平台游戏实现、跳跃手感、游戏AI和游戏引擎架构，也包括编程模式、游戏体验设计以及Unity官方手册和脚本API等内容",
            False,
        ),
        ("[1-9]", True),
        (
            "，这些内容可以为本项目的玩家控制、场景搭建、动画切换和敌人行为设计提供参考。",
            False,
        ),
    ],
)

p = doc.paragraphs[52]
set_body_paragraph_with_citation(
    p,
    [
        (
            "在游戏AI的设计中，有限状态机、行为树、分层状态机等都是比较常见的组织方式，拿毕业设计规模的2D动作游戏Demo来说，有限状态机的实现成本低、结构直观、调试时容易看出对象当前处在哪个阶段；普通敌人用待机、巡逻、追踪、攻击、受击、死亡等状态来描述行为，Boss可以在普通敌人基础上加上传送、施法、二阶段转场等状态来表现更复杂的战斗节奏；与把全部逻辑堆在一个Update函数中的方式相比，状态机可以让每一个状态只负责本阶段要做的事情，后续维护会轻一些",
            False,
        ),
        ("[10-11]", True),
        ("。", False),
    ],
)

p = doc.paragraphs[53]
set_body_paragraph_with_citation(
    p,
    [
        (
            "手感一般同输入响应速度、移动加速度、跳跃高度、碰撞检测稳定度、攻击前后摇、受击反馈等有着联系，本项目并没有试图构建起一个商业作品复杂的程度，而把重心放在了毕业Demo可以展现的核心体验部分上，即玩家能够较为流畅地移动、跳跃、冲刺、攻击，敌人可以在一定的范围内找到玩家并做出反应，Boss会给出阶段变化以及明显的攻击提示；这样的取舍有利于在较短的开发周期内得到稳定的、可以答辩的系统成果",
            False,
        ),
        ("[12-14]", True),
        ("。", False),
    ],
)

p = doc.paragraphs[98]
set_body_paragraph_with_citation(
    p,
    [
        ("本设计使用Unity引擎和C#语言进行开发，其主要的开发环境和工具见表3.1", False),
        ("[8-9,15]", True),
        ("。", False),
    ],
)

p = doc.paragraphs[101]
base = p.text.replace("[8-9]", "")
set_body_paragraph_with_citation(p, [(base.rstrip("。"), False), ("[8-9]", True), ("。", False)])

p = doc.paragraphs[217]
base = p.text.replace("[3,10-11]", "")
set_body_paragraph_with_citation(p, [(base.rstrip("。"), False), ("[3,10-11]", True), ("。", False)])

# 修正第 2 条会议论文，按 DiGRA 官方页面的作者顺序和题名标点整理。
set_reference_text(
    doc.paragraphs[309],
    "[2] Fasterholdt M, Pichlmair M, Holmgård C. You Say Jump, I Say How High? Operationalising the Game Feel of Jumping[C]//Proceedings of DiGRA/FDG 2016 Conference. DiGRA, 2016.",
)

doc.save(OUT)
print(f"backup={BACKUP}")
print(f"out={OUT}")
