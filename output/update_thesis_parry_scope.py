from pathlib import Path

from docx import Document


THESIS = Path(r"E:\Project_Metroidvania\Metroidvania\葛广一\毕业设计论文.docx")
FALLBACK = Path(r"E:\Project_Metroidvania\Metroidvania\葛广一\毕业设计论文_弹反适用范围补充版.docx")


def replace_paragraph(paragraphs, old, new):
    for paragraph in paragraphs:
        if paragraph.text == old:
            paragraph.text = new
            return
    raise RuntimeError(f"Paragraph not found: {old[:80]}")


def set_cell_text(cell, text):
    cell.text = text


doc = Document(THESIS)

replace_paragraph(
    doc.paragraphs,
    "本文围绕Unity平台下的2D横版动作游戏Demo展开，主要目标是完成一款类恶魔城风格的小型可运行作品，项目使用C#脚本进行功能编写，并通过有限状态机组织玩家、普通敌人和Boss的行为流程；角色移动、跳跃、冲刺、普通攻击和技能释放，主要依靠Rigidbody2D、碰撞检测以及Animator之间的配合来完成，战斗部分包含生命值管理、伤害判定、受击反馈、飞剑、黑洞和Boss二阶段激光等内容，关卡部分加入了落石机关、骷髅王和Tilemap通路事件，玩家会先经历基础战斗，再过渡到最终Boss战，系统还配有开始、暂停、死亡、通关菜单以及血条、音效和Windows打包支持，测试情况说明，该Demo能够完成开始、探索、战斗、机关躲避、Boss挑战和结算这一整条流程，可以作为毕业设计展示和论文分析的实践基础。",
    "本文围绕Unity平台下的2D横版动作游戏Demo展开，主要目标是完成一款类恶魔城风格的小型可运行作品，项目使用C#脚本进行功能编写，并通过有限状态机组织玩家、普通敌人和Boss的行为流程；角色移动、跳跃、冲刺、普通攻击、弹反和技能释放，主要依靠Rigidbody2D、碰撞检测以及Animator之间的配合来完成，战斗部分包含生命值管理、伤害判定、受击反馈、飞剑、黑洞和Boss二阶段激光等内容，关卡部分加入了落石机关、骷髅王和Tilemap通路事件，玩家会先经历基础战斗，再过渡到最终Boss战，系统还配有开始、暂停、死亡、通关菜单以及血条、音效和Windows打包支持，测试情况说明，该Demo能够完成开始、探索、战斗、机关躲避、Boss挑战和结算这一整条流程，可以作为毕业设计展示和论文分析的实践基础。",
)

replace_paragraph(
    doc.paragraphs,
    "玩家角色是系统交互的核心，角色需要能根据键盘输入完成左右移动、跳跃、冲刺、墙面滑落、墙跳和普通攻击，并在动画状态上表现出对应动作；为了减少输入和动画之间互相干扰，玩家行为不适合全部堆在一个脚本中，而要通过状态机拆分为待机、移动、跳跃、空中、冲刺、攻击、瞄准飞剑、接剑等状态。",
    "玩家角色是系统交互的核心，角色需要能根据键盘输入完成左右移动、跳跃、冲刺、墙面滑落、墙跳、普通攻击和Q键弹反，并在动画状态上表现出对应动作；为了减少输入和动画之间互相干扰，玩家行为不适合全部堆在一个脚本中，而要通过状态机拆分为待机、移动、跳跃、空中、冲刺、攻击、弹反、瞄准飞剑、接剑等状态。",
)

replace_paragraph(
    doc.paragraphs,
    "战斗系统既要支持玩家通过普通攻击、飞剑和黑洞对敌人造成伤害，也要支持敌人和Boss对玩家造成伤害，伤害不应只表现为数值变化，还要联动受击反馈、击退方向、无敌时间、死亡判定和音效播放；对于毕业Demo来说，战斗系统的重点在于判定要清楚、反馈要看得出来、流程也要尽量稳定。",
    "战斗系统既要支持玩家通过普通攻击、飞剑和黑洞对敌人造成伤害，也要支持敌人和Boss对玩家造成伤害，伤害不应只表现为数值变化，还要联动受击反馈、击退方向、无敌时间、死亡判定和音效播放；Q键弹反在当前Demo中与普通攻击同属基础近战交互，但它只针对开放反制窗口的骷髅类敌人，Boss攻击不纳入该弹反范围。对于毕业Demo来说，战斗系统的重点在于判定要清楚、反馈要看得出来、流程也要尽量稳定。",
)

replace_paragraph(
    doc.paragraphs,
    "普通攻击与技能释放还需要避免只看输入而忽略表现节奏。项目把部分判定放到动画关键帧上，玩家看到攻击动作挥出时，攻击区域才会真正检测敌人；飞剑则先进入瞄准状态，再在投掷事件触发后生成对象。菜单按钮点击后，系统会使用未受Time.timeScale影响的时间短暂屏蔽动作输入，这个处理专门解决开始游戏时鼠标点击被误当成攻击的问题。对动作游戏Demo而言，这类细节不显眼，却会直接影响操作是否顺手。",
    "普通攻击与技能释放还需要避免只看输入而忽略表现节奏。项目把部分判定放到动画关键帧上，玩家看到攻击动作挥出时，攻击区域才会真正检测敌人；Q键弹反也放在玩家基础战斗输入中处理，只有普通骷髅攻击动画开放反制窗口后，CounterAttack状态才会命中并造成反击伤害，Boss战仍主要依靠移动、冲刺和技能应对。飞剑则先进入瞄准状态，再在投掷事件触发后生成对象。菜单按钮点击后，系统会使用未受Time.timeScale影响的时间短暂屏蔽动作输入，这个处理专门解决开始游戏时鼠标点击被误当成攻击的问题。对动作游戏Demo而言，这类细节不显眼，却会直接影响操作是否顺手。",
)

replace_paragraph(
    doc.paragraphs,
    "界面提示方面，Demo在玩家初始区域上方的背景处放置操作提示，提示内容覆盖移动、跳跃、冲刺、普通攻击、弹反、飞剑和黑洞按键，并在开局一段时间后淡出；这种放置方式不会把文字压在角色身上，也能让首次试玩者先读到核心操作。敌方世界空间血条也根据试玩效果加宽加厚，普通骷髅、骷髅王和Boss受击时的生命变化更容易观察。",
    "界面提示方面，Demo在玩家初始区域上方的背景处放置操作提示，提示内容覆盖移动、跳跃、冲刺、普通攻击、对骷髅使用的Q键弹反、飞剑和黑洞按键，并在开局一段时间后淡出；这种放置方式不会把文字压在角色身上，也能让首次试玩者先读到核心操作。敌方世界空间血条也根据试玩效果加宽加厚，普通骷髅、骷髅王和Boss受击时的生命变化更容易观察。",
)

replace_paragraph(
    doc.paragraphs,
    "功能测试覆盖从游戏启动到Boss击败的完整路径，测试人员从开始菜单进入游戏后，先确认玩家不会因为点击按钮而立即攻击，随后测试移动、跳跃、冲刺、墙跳和普通攻击，再与骷髅敌人交互，观察敌人是否能发现玩家、追踪、攻击、受击和死亡，最后进入Boss区域，检查Boss一阶段近战、法术、传送、二阶段激光、死亡和通关菜单，这一路径基本覆盖Demo答辩时需要展示的主要内容。",
    "功能测试覆盖从游戏启动到Boss击败的完整路径，测试人员从开始菜单进入游戏后，先确认玩家不会因为点击按钮而立即攻击，随后测试移动、跳跃、冲刺、墙跳、普通攻击和Q键弹反，再与骷髅敌人交互，观察敌人是否能发现玩家、追踪、攻击、受击和死亡，同时确认弹反只在骷髅反制窗口内生效；最后进入Boss区域，检查Boss一阶段近战、法术、传送、二阶段激光、死亡和通关菜单，并确认Boss攻击不走普通骷髅的弹反流程，这一路径基本覆盖Demo答辩时需要展示的主要内容。",
)

replace_paragraph(
    doc.paragraphs,
    "在玩家控制部分，系统实现移动、跳跃、冲刺、墙面交互、普通攻击、飞剑和黑洞技能等功能；在敌人AI部分，骷髅敌人能够完成巡逻、发现玩家、追踪、攻击、受击到死亡的流程；在关卡交互部分，落石机关能够根据玩家位置独立触发并循环回收，骷髅王被击败后能够解除Tilemap通路；在Boss设计部分，DeathBringer具备近战攻击、法术攻击、传送、二阶段激光转场和死亡反馈；在UI部分，系统实现开始、暂停、死亡和通关菜单，并通过血条显示角色生命值，从测试结果看，Demo能够在Windows环境下独立运行并完成主要战斗与关卡演示流程。",
    "在玩家控制部分，系统实现移动、跳跃、冲刺、墙面交互、普通攻击、骷髅弹反、飞剑和黑洞技能等功能；在敌人AI部分，骷髅敌人能够完成巡逻、发现玩家、追踪、攻击、受击到死亡的流程；在关卡交互部分，落石机关能够根据玩家位置独立触发并循环回收，骷髅王被击败后能够解除Tilemap通路；在Boss设计部分，DeathBringer具备近战攻击、法术攻击、传送、二阶段激光转场和死亡反馈，Boss攻击未纳入当前弹反范围；在UI部分，系统实现开始、暂停、死亡和通关菜单，并通过血条显示角色生命值，从测试结果看，Demo能够在Windows环境下独立运行并完成主要战斗与关卡演示流程。",
)

system_requirements = doc.tables[0]
set_cell_text(system_requirements.rows[1].cells[1], "移动、跳跃、冲刺、墙面交互、普通攻击、弹反")
set_cell_text(system_requirements.rows[1].cells[2], "保证角色能够完成横版动作游戏的基础操作与近战反制")

player_states = doc.tables[4]
set_cell_text(player_states.rows[4].cells[0], "Attack / CounterAttack")
set_cell_text(player_states.rows[4].cells[1], "Attack、CounterAttack")
set_cell_text(player_states.rows[4].cells[2], "播放普通攻击动画，并在骷髅开放窗口时处理Q键弹反判定")

function_tests = doc.tables[9]
set_cell_text(function_tests.rows[3].cells[0], "普通攻击与弹反")
set_cell_text(function_tests.rows[3].cells[1], "靠近敌人执行普通攻击，并在骷髅开放反制窗口时按Q")
set_cell_text(function_tests.rows[3].cells[2], "普通攻击产生受击和击退；骷髅可被弹反，Boss不进入该弹反流程")

try:
    doc.save(THESIS)
    print(f"Updated thesis parry scope descriptions: {THESIS}")
except PermissionError:
    doc.save(FALLBACK)
    print(f"Main thesis is locked; wrote fallback: {FALLBACK}")
