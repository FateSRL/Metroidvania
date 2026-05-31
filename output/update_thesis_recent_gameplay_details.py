# -*- coding: utf-8 -*-
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX_PATH = os.environ["THESIS_DOCX"]

INSERTIONS = {
    250: [
        "技能和Boss联动回归测试还需要检查临界时机：飞剑接回前不应提前出现冷却，黑洞结束后要在玩家落地时开始计时，Boss处于攻击、施法或传送动作时被黑洞覆盖仍应实际扣血；同时还要观察黑洞最后一段伤害击杀普通骷髅时，死亡动画和碰撞释放是否连续执行。把这些检查放进主流程回归后，可以减少局部修复只在单一场景下有效的问题。"
    ],
    228: [
        "界面提示方面，Demo在玩家初始区域上方的背景处放置操作提示，提示内容覆盖移动、跳跃、冲刺、普通攻击、弹反、飞剑和黑洞按键，并在开局一段时间后淡出；这种放置方式不会把文字压在角色身上，也能让首次试玩者先读到核心操作。敌方世界空间血条也根据试玩效果加宽加厚，普通骷髅、骷髅王和Boss受击时的生命变化更容易观察。"
    ],
    223: [
        "Boss与黑洞的交互也做了单独约束。攻击、施法和传送动作如果完全保持原状态，黑洞命中会因为传送无敌或状态占用而显得收益不足；当前实现中，Boss被黑洞捕获时会打断这三类动作，回到原地停住的Idle，再由公共伤害流程扣除生命值并给出受击闪白反馈。二阶段激光转场仍保留无敌保护，避免阶段演出被范围技能截断。结合试玩节奏，Boss基础移动速度也做了小幅提高，使追踪阶段的压迫感更明显。"
    ],
    173: [
        "技能冷却在本项目中按交互结果计算，而不是在按键触发后统一开始计时。飞剑投出后仍可能停留在敌人或地面上，只有玩家将其召回并接住，冷却才会启动；黑洞则在范围效果结束后等待玩家重新落地，再进入冷却。这样处理后，技能动作、状态切换和界面提示能够对齐，技能处于冷却期时也不会先播放上升动作再被脚本拦回。冷却显示由UI_SkillCooldownPanel读取技能剩余比例，在屏幕右上角通过图标和径向遮罩提示飞剑、黑洞是否可再次使用。",
        "黑洞多段伤害还补充了死亡边界处理。控制器只有在实际命中敌人后才消耗一次伤害次数，当范围内目标仍在时，会等待既定次数完成后再结束技能；若最后一段伤害正好击杀敌人，循环会先保存当前Enemy引用，再继续处理冻结与释放，避免碰撞体关闭后敌人列表变化引发索引异常。被冻结目标进入死亡流程时还要恢复Animator速度，否则死亡参数虽然已经写入，动画也可能停留在冻结帧。"
    ],
}


def insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = paragraph.style
    new_paragraph.add_run(text)
    return new_paragraph


def main() -> None:
    document = Document(DOCX_PATH)
    existing_text = {paragraph.text for paragraph in document.paragraphs}
    inserted = 0

    for index in sorted(INSERTIONS, reverse=True):
        anchor = document.paragraphs[index]
        current = anchor

        for text in INSERTIONS[index]:
            if text in existing_text:
                continue

            current = insert_after(current, text)
            inserted += 1

    document.save(DOCX_PATH)
    print(f"inserted_paragraphs={inserted}")
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
