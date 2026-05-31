from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SRC = Path("output/thesis_final_format_source.docx")
TEMPLATE = Path("output/thesis_template.docx")
OUT = Path("output/thesis_final_format_adjusted.docx")


BODY = "af7"
CHAPTER = "ac"
SECTION = "af0"
SUBSECTION = "a3"
CN_ABSTRACT = "af5"
EN_ABSTRACT = "af9"
REFERENCE = "afa"
PICTURE = "afe"
FORMULA = "afc"
TOC_IDS = {"TOC1", "TOC2", "TOC3", "TOC", "TOCHeading"}


def style_by_id(doc, style_id):
    for style in doc.styles:
        if style.style_id == style_id:
            return style
    return None


def set_style_font(style, east_asia=None, latin=None, size=None, bold=None):
    if style is None:
        return
    font = style.font
    if latin:
        font.name = latin
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    font.color.rgb = RGBColor(0, 0, 0)

    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = deepcopy(rPr._new_rFonts())
        rPr.append(rFonts)
    if east_asia:
        rFonts.set(qn("w:eastAsia"), east_asia)
    if latin:
        rFonts.set(qn("w:ascii"), latin)
        rFonts.set(qn("w:hAnsi"), latin)


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=12, bold=None):
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.underline = False
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = deepcopy(rPr._new_rFonts())
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)


def set_para_format(paragraph, alignment=None, first_indent=None, before=None, after=None, line=None, rule=None):
    fmt = paragraph.paragraph_format
    if alignment is not None:
        fmt.alignment = alignment
    if first_indent is not None:
        fmt.first_line_indent = Pt(first_indent)
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing = line if not isinstance(line, (int, float)) else line
    if rule is not None:
        fmt.line_spacing_rule = rule


def copy_section_layout(src_doc, template_doc):
    src_sections = src_doc.sections
    template_section = template_doc.sections[0]
    for sec in src_sections:
        sec.top_margin = template_section.top_margin
        sec.bottom_margin = template_section.bottom_margin
        sec.left_margin = template_section.left_margin
        sec.right_margin = template_section.right_margin
        sec.page_width = template_section.page_width
        sec.page_height = template_section.page_height


def is_caption(text):
    t = text.strip()
    return (
        len(t) >= 3
        and (t[0] in {"图", "表"})
        and t[1].isdigit()
        and ("." in t[:6] or " " in t[:6])
    )


def is_reference_heading(text):
    t = text.strip().replace(" ", "")
    return t in {"参考文献", "致谢", "结论"} or t.startswith("附录")


def format_target_styles(doc):
    set_style_font(style_by_id(doc, BODY), "宋体", "Times New Roman", 12, False)
    set_style_font(style_by_id(doc, CHAPTER), "黑体", "Times New Roman", 15, True)
    set_style_font(style_by_id(doc, SECTION), "黑体", "Times New Roman", 14, True)
    set_style_font(style_by_id(doc, SUBSECTION), "黑体", "Times New Roman", 12, True)
    set_style_font(style_by_id(doc, CN_ABSTRACT), "黑体", "Times New Roman", 15, True)
    set_style_font(style_by_id(doc, EN_ABSTRACT), "Times New Roman", "Times New Roman", 16, True)
    set_style_font(style_by_id(doc, REFERENCE), "宋体", "Times New Roman", 10.5, False)
    set_style_font(style_by_id(doc, PICTURE), "宋体", "Times New Roman", 10.5, False)
    set_style_font(style_by_id(doc, FORMULA), "宋体", "Times New Roman", 12, False)

    style_format = {
        CHAPTER: (WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 5, 1.5, WD_LINE_SPACING.ONE_POINT_FIVE),
        SECTION: (WD_ALIGN_PARAGRAPH.LEFT, 0, 6, 0, 1.5, WD_LINE_SPACING.ONE_POINT_FIVE),
        SUBSECTION: (WD_ALIGN_PARAGRAPH.LEFT, 0, 2.5, 0, 1.5, WD_LINE_SPACING.ONE_POINT_FIVE),
        BODY: (WD_ALIGN_PARAGRAPH.JUSTIFY, 24, 0, 0, 1.25, WD_LINE_SPACING.MULTIPLE),
        REFERENCE: (WD_ALIGN_PARAGRAPH.JUSTIFY, None, 0, 0, Pt(17), WD_LINE_SPACING.AT_LEAST),
    }
    for sid, values in style_format.items():
        style = style_by_id(doc, sid)
        if style is None:
            continue
        align, first, before, after, line, rule = values
        pf = style.paragraph_format
        pf.alignment = align
        if first is not None:
            pf.first_line_indent = Pt(first)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line
        pf.line_spacing_rule = rule


def format_paragraph(paragraph, in_english_abstract=False):
    sid = paragraph.style.style_id
    text = paragraph.text.strip()
    if sid in TOC_IDS or sid.startswith("TOC"):
        return

    if sid == CHAPTER or is_reference_heading(text):
        set_para_format(paragraph, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 5, 1.5, WD_LINE_SPACING.ONE_POINT_FIVE)
        for run in paragraph.runs:
            set_run_font(run, "黑体", "Times New Roman", 15, True)
        return

    if sid == SECTION:
        set_para_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT, 0, 6, 0, 1.5, WD_LINE_SPACING.ONE_POINT_FIVE)
        for run in paragraph.runs:
            set_run_font(run, "黑体", "Times New Roman", 14, True)
        return

    if sid == SUBSECTION:
        set_para_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT, 0, 2.5, 0, 1.5, WD_LINE_SPACING.ONE_POINT_FIVE)
        for run in paragraph.runs:
            set_run_font(run, "黑体", "Times New Roman", 12, True)
        return

    if sid == CN_ABSTRACT:
        set_para_format(paragraph, WD_ALIGN_PARAGRAPH.CENTER, 0, 10, 14, None, None)
        for run in paragraph.runs:
            set_run_font(run, "黑体", "Times New Roman", 15, True)
        return

    if sid == EN_ABSTRACT:
        set_para_format(paragraph, WD_ALIGN_PARAGRAPH.CENTER, 0, 10, 14, None, None)
        for run in paragraph.runs:
            set_run_font(run, "Times New Roman", "Times New Roman", 16, True)
        return

    if sid == REFERENCE:
        fmt = paragraph.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.left_indent = Pt(21)
        fmt.first_line_indent = Pt(-21)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(17)
        fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        for run in paragraph.runs:
            set_run_font(run, "宋体", "Times New Roman", 10.5, False)
        return

    if is_caption(text):
        set_para_format(paragraph, WD_ALIGN_PARAGRAPH.CENTER, 0, 3, 3, 1.25, WD_LINE_SPACING.MULTIPLE)
        for run in paragraph.runs:
            set_run_font(run, "宋体", "Times New Roman", 12, True)
        return

    if sid == BODY:
        first_indent = 0 if text.startswith(("关键词", "Key words", "Keywords")) or in_english_abstract else 24
        line = 1.5 if in_english_abstract and not text.startswith(("Key words", "Keywords")) else 1.25
        set_para_format(paragraph, WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent, 0, 0, line, WD_LINE_SPACING.MULTIPLE)
        east = "Times New Roman" if in_english_abstract else "宋体"
        for run in paragraph.runs:
            set_run_font(run, east, "Times New Roman", 12, None)
        return


def format_tables(doc):
    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.style.style_id in TOC_IDS:
                        continue
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.25
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                    if row_index == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(run, "宋体", "Times New Roman", 10.5, row_index == 0)


def main():
    doc = Document(SRC)
    template = Document(TEMPLATE)
    copy_section_layout(doc, template)
    format_target_styles(doc)

    in_english_abstract = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "ABSTRACT":
            in_english_abstract = True
            format_paragraph(paragraph)
            continue
        if text.startswith("1 ") and paragraph.style.style_id == CHAPTER:
            in_english_abstract = False
        format_paragraph(paragraph, in_english_abstract=in_english_abstract)

    format_tables(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
