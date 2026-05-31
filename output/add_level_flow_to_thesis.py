# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path(r"E:\Project_Metroidvania\Metroidvania\output\thesis_final_current.docx")
OUT = Path(r"E:\Project_Metroidvania\Metroidvania\output\thesis_final_with_level_flow.docx")

BODY_STYLE = "论文正文"


def iter_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph


def find_paragraph(doc, startswith=None, contains=None, exact=None):
    for paragraph in iter_paragraphs(doc):
        text = paragraph.text.strip()
        if exact is not None and text == exact:
            return paragraph
        if startswith is not None and text.startswith(startswith):
            return paragraph
        if contains is not None and contains in text:
            return paragraph
    raise ValueError(f"anchor not found: startswith={startswith!r}, contains={contains!r}, exact={exact!r}")


def insert_after(paragraph, text, style=BODY_STYLE):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def replace_paragraph(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


doc = Document(SOURCE)

insertions = [
    (
        {"startswith": "为了使Demo不直接进入Boss战"},
        "从玩家实际经过的顺序看，完整流程被设计成“右侧尝试受阻、左侧完成挑战、开门后再进入平台跳跃、最后进入Boss战”的结构。玩家出生后先向右与普通小兵战斗，继续前进时会发现通往Boss区域的门暂时无法通过；随后玩家需要转向左侧落石区域，利用冲刺躲开从上方落下的石块，并击败更高生命值的骷髅王。骷髅王死亡后，门前阻挡砖块被移除，玩家再回到右侧，经过一段需要蹬墙反跳和跳上石柱的平台路线，最后从高处下落进入Boss区域。"
    ),
    (
        {"startswith": "Demo采用线性的演示流程"},
        "在场景空间上，门控区域并不是单纯的墙体阻挡，它承担了路线引导的作用。玩家向右探索时先接触普通骷髅和被封闭的通路，知道Boss区域暂时不能直接进入；左侧落石和骷髅王区域构成支线挑战，玩家完成这段战斗后再回到右侧。通路开启之后，平台、墙面和右侧石柱组成第二段移动考验，玩家需要连续两次蹬墙反跳上升，再跳到右侧石柱，最后下落进入Boss区域。这样的布局把战斗、机关、平台移动和Boss战串成一个较完整的演示路线。"
    ),
    (
        {"startswith": "关卡流程以基础操作展示"},
        "按照实际游玩顺序，关卡事件可以拆成六个节点：第一，玩家出生后向右移动并与普通小兵交战；第二，玩家继续向右探索，遇到暂时无法通过的门，形成明确的路线阻挡；第三，玩家转向左侧落石区，通过单块石头触发的下落机关；第四，玩家击败骷髅王，触发通路砖块消失；第五，玩家回到门后区域，通过墙面反跳和石柱跳跃完成平台段；第六，玩家从平台右侧下落到Boss区域，进入最终战斗。这里没有使用文字任务提示，而是通过空间阻挡、敌人配置和可通行区域变化来引导玩家。"
    ),
    (
        {"startswith": "就目前的Demo而言，机关和精英怪一起填补了Boss前面的过渡区"},
        "开门之后的跳跃段主要用于验证玩家控制系统是否真正能服务于关卡。该区域不是简单摆放几个平台，而是让玩家先贴墙滑落，再通过两次墙面反跳爬升，随后跳到右侧石柱并从石柱处下落到Boss区域；这一段会同时用到WallCheck、GroundCheck、PlayerWallSlideState和PlayerWallJumpState等逻辑。如果墙体碰撞、相机跟随或者平台高度配置不合理，玩家就会在这里卡住，所以它在流程上连接机关和Boss，在实现上也检验了玩家移动模块的稳定性。"
    ),
    (
        {"startswith": "功能测试不能只停留在场景中单个对象上点开查看"},
        "完整路线测试时，测试人员需要按照实际演示顺序重新走一遍：从开始菜单进入游戏，先向右击败普通小兵，确认前方通路仍被门阻挡；随后折返回左侧落石区，检查每个石块只在玩家经过其正下方时下落，并通过冲刺躲避；击败骷髅王后观察阻挡砖块是否消失，再回到右侧完成连续蹬墙反跳、跳上右侧石柱和下落进入Boss区域的流程。该测试不是单纯检查某一个脚本，而是检查普通敌人、落石、骷髅王、Tilemap门、墙跳平台和Boss触发区域能否连续衔接。"
    ),
]

for anchor, text in insertions:
    insert_after(find_paragraph(doc, **anchor), text)

replace_paragraph(
    find_paragraph(doc, startswith="从测试覆盖角度来说，目前Demo的验证重点已经从单个功能是否可以运行"),
    "从测试覆盖角度来说，目前Demo的验证重点已经从单个功能是否可以运行，转向了完整路线能否连续运行。玩家可以从开始菜单进入游戏，先向右和普通小兵战斗，再发现前方门暂时无法通过；随后玩家折返左侧，躲避落石并击败骷髅王，通路开启后再经过需要连续蹬墙反跳的平台区，跳上右侧石柱并下落到Boss区域，最终完成Boss战和通关菜单。该路径包含输入、物理、动画、AI、Tilemap事件、UI菜单、音效等多个模块，如果这条路线可以稳定运行，那么项目就具备了比较完整的毕业演示价值。"
)

replace_paragraph(
    find_paragraph(doc, startswith="对于答辩展示来说，系统说明要对应到现场操作"),
    "对于答辩展示来说，系统说明要对应到现场操作，目前的Demo可以按照固定的路线进行演示：从开始菜单进入游戏后，先展示玩家移动、跳跃、冲刺、普通攻击和Q键弹反，再向右与普通骷髅交互，说明敌人检测、追踪、攻击、受击和死亡流程；接着让玩家发现前方通路无法直接通过，转向左侧落石机关，说明触发器、协程、即死判定和单块检测；击败骷髅王后展示Tilemap通路开启，再完成两次蹬墙反跳、跳上右侧石柱并下落进入Boss区域，最后展示Boss一阶段攻击、传送、法术、二阶段激光和通关菜单。这条演示路线与论文结构基本一致，可以把抽象模块落到具体画面中。"
)

doc.save(OUT)
print(f"wrote={OUT}")
print("insertions=5")
print("replacements=2")
